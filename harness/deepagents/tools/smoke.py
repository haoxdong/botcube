from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import dateutil
import docx
import img2pdf
import joblib
import matplotlib
import mpmath
import numpy
import openpyxl
import pandas
import pdf2image
import pdfplumber
import PIL
import pptx
import pyarrow
import pypdf
import pypdfium2
import pytz
import reportlab
import requests
import scipy
import seaborn
import sklearn
import statsmodels
import sympy
import tqdm
import tzdata
import xlrd
import xlsxwriter
from matplotlib.font_manager import FontProperties
from PIL import Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

matplotlib.use('Agg')
from matplotlib import pyplot


def main(output_dir: Path) -> None:
    modules = (pandas, numpy, scipy, sklearn, statsmodels, matplotlib, seaborn,
               pyarrow, openpyxl, xlsxwriter, xlrd, PIL, pptx, docx, pypdf,
               pdfplumber, pypdfium2, pdf2image, reportlab, img2pdf, sympy,
               mpmath, tqdm, dateutil, pytz, tzdata, joblib, requests)
    print('Imported ' + ', '.join(module.__name__ for module in modules))
    for executable in ('unzip', '7z', 'bc', 'rg', 'fd', 'sqlite3', 'pdftoppm', 'pdftotext', 'fc-match'):
        subprocess.run(['bash', '-c', 'command -v "$1"', 'smoke', executable], check=True, capture_output=True)
    font = subprocess.run(['fc-match', '-f', '%{file}', 'Noto Sans'], check=True, capture_output=True, text=True).stdout
    assert 'NotoSans' in font, font
    for family in ('Noto Sans CJK SC', 'Noto Color Emoji'):
        match = subprocess.run(['fc-match', '-f', '%{family}', family], check=True, capture_output=True, text=True).stdout
        assert family in match, match
    pdfmetrics.registerFont(TTFont('Noto', font))
    output_dir.mkdir(parents=True, exist_ok=True)
    presentation = pptx.Presentation()
    for title in ('Iran oil', 'Supply', 'Outlook'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        assert slide.shapes.title is not None
        slide.shapes.title.text = title
    presentation.save(str(output_dir / 'slides.pptx'))
    reopened = pptx.Presentation(str(output_dir / 'slides.pptx'))
    titles: list[str] = []
    for slide in reopened.slides:
        assert slide.shapes.title is not None
        titles.append(slide.shapes.title.text)
    assert titles == ['Iran oil', 'Supply', 'Outlook']
    workbook = openpyxl.Workbook()
    assert workbook.active is not None
    workbook.active['A1'] = 'Iran oil'
    workbook.save(output_dir / 'data.xlsx')
    reopened_workbook = openpyxl.load_workbook(output_dir / 'data.xlsx')
    assert reopened_workbook.active is not None
    assert reopened_workbook.active['A1'].value == 'Iran oil'
    document = docx.Document()
    document.add_paragraph('Iran oil')
    document.save(str(output_dir / 'report.docx'))
    assert docx.Document(str(output_dir / 'report.docx')).paragraphs[0].text == 'Iran oil'
    pdf_path = output_dir / 'report.pdf'
    canvas = Canvas(str(pdf_path))
    canvas.setFont('Noto', 14)
    canvas.drawString(72, 720, 'Iran oil · Ελληνικά · Русский')
    canvas.save()
    assert 'Ελληνικά' in pypdf.PdfReader(pdf_path).pages[0].extract_text()
    pages = pdf2image.convert_from_path(pdf_path, dpi=36)
    assert len(pages) == 1 and pages[0].width > 0
    pyplot.plot([1, 2], [3, 4])
    pyplot.title('Oil · Ελληνικά', fontproperties=FontProperties(fname=font))
    pyplot.savefig(output_dir / 'chart.png')
    pyplot.close()
    with Image.open(output_dir / 'chart.png') as image:
        assert image.format == 'PNG' and image.width > 0 and image.height > 0
    for path in output_dir.iterdir():
        assert path.stat().st_size > 0, path
    print('Roundtripped PPTX, XLSX, DOCX, PDF and PNG; Poppler and Noto fonts verified')


if __name__ == '__main__':
    main(Path(sys.argv[1]))
